#!/usr/bin/env python3
"""从 data.xlsx 的「族谱」表生成 data.js。

用法：python build_data.py [xlsx路径] [输出路径]
需要：pip install openpyxl
只读取 A 到 C 列（编号、姓名、父亲编号）；其他列（备注等）不会进入网站。
数据有错误时会列出所有问题并退出（退出码 1），不会覆盖 data.js。
"""
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "data.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "data.js"
SHEET = "族谱"

HEADER = """\
// 黄氏族谱数据 —— 由 build_data.py 从 data.xlsx 自动生成，请不要手动修改。
// 更新方法：只需修改 Excel（data.xlsx），上传到 GitHub 后会自动重新生成本文件。
// 格式：编号  姓名  父亲编号（Tab 分隔）。始祖的父亲编号留空。

const FAMILY_DATA = `
"""


def as_id(value, where, errors):
    """把单元格转成整数编号的字符串；空格返回 ''。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not re.fullmatch(r"\d+", text):
        errors.append(f"{where}：编号“{text}”不是数字")
        return ""
    return str(int(text))


def read_people(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"错误：{path.name} 里找不到名为「{SHEET}」的工作表")
    ws = wb[SHEET]

    errors, people = [], []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, max_col=3, values_only=True), start=2):
        raw_id, raw_name, raw_parent = row
        name = "" if raw_name is None else str(raw_name).strip()
        if raw_id is None and not name and raw_parent is None:
            continue  # 空行
        where = f"第 {row_no} 行"
        n_before = len(errors)
        pid = as_id(raw_id, where, errors)
        parent = as_id(raw_parent, where + "的父亲", errors)
        if not pid:
            if len(errors) == n_before:
                errors.append(f"{where}：缺少编号（姓名：{name or '空'}）")
            continue
        if not name:
            errors.append(f"{where}：编号 {pid} 没有姓名")
            continue
        if re.search(r"[\s,，`\\]|\$\{", name):
            errors.append(f"{where}：姓名“{name}”含有空格、逗号或特殊符号")
            continue
        people.append((row_no, pid, name, parent))
    return people, errors


def validate(people, errors):
    seen = {}
    for row_no, pid, name, _ in people:
        if pid in seen:
            errors.append(f"第 {row_no} 行：编号 {pid}（{name}）与第 {seen[pid]} 行重复")
        else:
            seen[pid] = row_no
    parent_of = {pid: parent for _, pid, _, parent in people}
    roots = [pid for pid, parent in parent_of.items() if not parent]
    for row_no, pid, name, parent in people:
        if parent == pid:
            errors.append(f"第 {row_no} 行：{name} 的父亲编号是自己")
        elif parent and parent not in parent_of:
            errors.append(f"第 {row_no} 行：{name} 的父亲编号 {parent} 不存在")
    # 循环检查（A 的父亲是 B，B 的父亲又是 A）
    for pid in parent_of:
        cur, steps = pid, 0
        while cur and cur in parent_of and steps <= len(parent_of):
            cur, steps = parent_of[cur], steps + 1
        if steps > len(parent_of):
            errors.append(f"编号 {pid} 的父亲关系形成了循环")
            break
    if len(roots) != 1:
        names = "、".join(f"{pid}" for pid in roots) or "无"
        errors.append(f"应该只有一位始祖（父亲编号留空），目前有 {len(roots)} 位（编号：{names}）")


def main():
    if not XLSX.exists():
        sys.exit(f"错误：找不到 {XLSX}")
    people, errors = read_people(XLSX)
    validate(people, errors)
    if errors:
        print(f"{XLSX.name} 有 {len(errors)} 个问题，data.js 没有更新：", file=sys.stderr)
        for e in errors:
            print("  - " + e, file=sys.stderr)
        sys.exit(1)

    lines = ["编号\t姓名\t父亲编号"] + [f"{pid}\t{name}\t{parent}" for _, pid, name, parent in people]
    content = HEADER + "\n".join(lines) + "\n`;\n"
    if OUT.exists() and OUT.read_text(encoding="utf-8") == content:
        print(f"{OUT.name} 已是最新（{len(people)} 人），无需更改")
        return
    OUT.write_text(content, encoding="utf-8", newline="\n")
    print(f"已生成 {OUT.name}：{len(people)} 人")


if __name__ == "__main__":
    main()
